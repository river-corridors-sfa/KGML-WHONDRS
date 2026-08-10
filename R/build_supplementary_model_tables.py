import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DERIVED = ROOT / "derived"
OUT_XLSX = DERIVED / "Supplementary_Table_model_performance.xlsx"
OUT_CSV = DERIVED / "Supplementary_Table_model_performance_primary.csv"


def read_csv(name):
    with (DERIVED / name).open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def numeric(value):
    if value in (None, "", "NA"):
        return None
    try:
        return int(value)
    except ValueError:
        try:
            return float(value)
        except ValueError:
            return value


def model_metadata(model):
    if model == "Mechanistic fixed Vh0":
        return "Mechanistic", "Mechanistic model", "None", "Observed log10(R)"
    family = "KGML" if model.startswith("KGML") else "Pure ML"
    algorithm = "GAM" if "GAM" in model else ("XGBoost" if "XGBoost" in model else "Random forest")
    predictor_set = next((x for x in ("Bulk+FTICR", "FTICR", "Bulk", "NPOC") if x in model), "NPOC")
    target = "Residual dR" if family == "KGML" else "Observed log10(R)"
    return family, algorithm, predictor_set, target


primary_raw = read_csv("kgml_cv_stats_all.csv")
primary = []
for row in primary_raw:
    family, algorithm, predictor_set, target = model_metadata(row["Model"])
    primary.append({
        "Validation": row["Validation"],
        "Model": row["Model"],
        "Model_family": family,
        "Algorithm": algorithm,
        "Predictor_set": predictor_set,
        "Training_target": target,
        "Vh0": None if family == "Pure ML" else numeric(row["Vh0"]),
        "n_predictions": numeric(row["n"]),
        "n_unique_samples": 141,
        "n_sites": 65,
        "RMSE": numeric(row["RMSE"]),
        "Pearson_r": numeric(row["Pearson_r"]),
        "CCC": numeric(row["CCC"]),
        "Bias": numeric(row["Bias"]),
        "R2": numeric(row["R2"]),
    })

primary.sort(key=lambda x: (x["Validation"], x["RMSE"]))

sensitivity = read_csv("clean_validation_sensitivity_stats.csv")
for row in sensitivity:
    if row["Validation"] == "Grouped 5-fold CV":
        row["Validation"] = "Site-grouped 5-fold CV"
    for key in ("Vh0", "n", "RMSE", "Pearson_r", "CCC", "Bias", "R2"):
        row[key] = numeric(row[key])

definitions = read_csv("clean_validation_predictor_manifest.csv")

metrics = [
    {"Metric": "RMSE", "Definition": "Root mean squared error between observed and predicted log10 respiration rate.", "Direction": "Lower is better"},
    {"Metric": "Pearson r", "Definition": "Pearson correlation between observed and predicted log10 respiration rate.", "Direction": "Higher is better"},
    {"Metric": "CCC", "Definition": "Lin's concordance correlation coefficient; measures agreement with the 1:1 line.", "Direction": "Higher is better"},
    {"Metric": "Bias", "Definition": "Mean prediction error (predicted minus observed) on the log10 scale.", "Direction": "Closer to zero is better"},
    {"Metric": "R2", "Definition": "1 minus residual sum of squares divided by total sum of squares on held-out predictions.", "Direction": "Higher is better; may be negative"},
]

notes = [
    ["Supplementary Table", "Model definitions and held-out performance for the WHONDRS S19S respiration models"],
    ["Primary performance", "Fixed-Vh0 model comparison under site-grouped 5-fold CV and leave-one-site-out CV (LOSO)."],
    ["Validation design", "All samples from a site remain together. Site-grouped 5-fold CV holds out groups of sites; LOSO holds out one complete site per iteration."],
    ["Repeated CV counts", "The site-grouped 5-fold analysis contains 1,410 held-out predictions because 141 samples were evaluated across 10 repeats. LOSO contains one held-out prediction per sample (n = 141)."],
    ["Sensitivity tab", "Full model statistics across predictor sets and fixed Vh0 values; NA Vh0 denotes pure-ML models for which Vh0 is not applicable."],
    ["Response scale", "Performance metrics are calculated for log10 respiration rate."],
    ["Source files", "kgml_cv_stats_all.csv, clean_validation_sensitivity_stats.csv, and clean_validation_predictor_manifest.csv"],
]


def write_rows(ws, rows):
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    return headers


combined = []
for row in primary:
    combined.append({
        "Analysis": "Primary performance",
        "Validation": row["Validation"],
        "Model": row["Model"],
        "Predictor_set": row["Predictor_set"],
        "Vh0": row["Vh0"],
        "RMSE": row["RMSE"],
        "Pearson_r": row["Pearson_r"],
        "CCC": row["CCC"],
    })

model_labels = {
    "KGML RF Bulk": "KGML-Bulk",
    "KGML RF Bulk+FTICR": "KGML-Bulk+FTICR",
    "KGML RF FTICR": "KGML-FTICR",
    "KGML RF NPOC": "KGML-NPOC",
    "Mechanistic": "Mechanistic fixed Vh0",
}
seen = set()
for row in sensitivity:
    if row["Vh0"] != 0.5:
        continue
    # The mechanistic prediction is identical across predictor-set labels.
    key = (row["Validation"], row["Model"])
    if row["Model"] == "Mechanistic" and key in seen:
        continue
    seen.add(key)
    combined.append({
        "Analysis": "Vh0 sensitivity (0.5)",
        "Validation": row["Validation"],
        "Model": model_labels[row["Model"]],
        "Predictor_set": "None" if row["Model"] == "Mechanistic" else row["Predictor_set"],
        "Vh0": 0.5,
        "RMSE": row["RMSE"],
        "Pearson_r": row["Pearson_r"],
        "CCC": row["CCC"],
    })

analysis_order = {
    "Primary performance": 1,
    "Vh0 sensitivity (0.5)": 2,
}
validation_order = {
    "Site-grouped 5-fold CV": 1,
    "Leave-one-site-out CV": 2,
}
predictor_order = {
    "None": 0,
    "NPOC": 1,
    "Bulk": 2,
    "FTICR": 3,
    "Bulk+FTICR": 4,
}


def model_order(model):
    if model.startswith("Mechanistic"):
        return 1
    if model == "KGML-NPOC":
        return 2
    if model == "KGML-GAM NPOC":
        return 3
    if model.startswith("KGML"):
        return 4
    if model.startswith("Pure RF"):
        return 5
    if model.startswith("Pure XGBoost"):
        return 6
    return 99


combined.sort(key=lambda x: (
    analysis_order[x["Analysis"]],
    validation_order[x["Validation"]],
    model_order(x["Model"]),
    predictor_order[x["Predictor_set"]],
))

wb = Workbook()
ws = wb.active
ws.title = "Model performance"
write_rows(ws, combined)

navy = "17365D"
blue = "D9EAF7"
white = "FFFFFF"
thin = Side(style="thin", color="B7C9D6")

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
ws.sheet_view.showGridLines = False
for cell in ws[1]:
    cell.fill = PatternFill("solid", fgColor=navy)
    cell.font = Font(color=white, bold=True)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    cell.border = Border(bottom=thin)
ws.row_dimensions[1].height = 30
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = Border(bottom=thin)
    if row[0].row % 2 == 0:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="F5F9FC")
for col_idx, column in enumerate(ws.columns, 1):
    values = [str(c.value or "") for c in column]
    width = min(max(max(map(len, values)) + 2, 11), 30)
    ws.column_dimensions[get_column_letter(col_idx)].width = width

headers = {cell.value: cell.column for cell in ws[1]}
for metric in ("Vh0", "RMSE", "Pearson_r", "CCC"):
    for cells in ws.iter_cols(min_col=headers[metric], max_col=headers[metric], min_row=2):
        for cell in cells:
            cell.number_format = "0.000"

wb.save(OUT_XLSX)

with OUT_CSV.open("w", newline="", encoding="utf-8") as handle:
    writer = csv.DictWriter(handle, fieldnames=list(combined[0].keys()))
    writer.writeheader()
    writer.writerows(combined)

# Read back the workbook to verify that all expected records were written.
check = load_workbook(OUT_XLSX, read_only=True, data_only=True)
assert check.sheetnames == ["Model performance"]
assert check["Model performance"].max_row == len(combined) + 1
print(f"Wrote {OUT_XLSX}")
print(f"Wrote {OUT_CSV}")
print(f"Combined rows: {len(combined)}")
